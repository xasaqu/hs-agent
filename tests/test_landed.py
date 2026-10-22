"""头程落地成本：美国税则解析、逐层关税、计费重和装箱校验、退税、多家报价对比、账单核对、报价导入。"""

import asyncio
import io
import json
import math
from datetime import date

import pytest
from openpyxl import load_workbook

from hs_agent.cn_tariff import CnTariff
from hs_agent.duty import Policy, entry_fees, line_duty
from hs_agent.landed import (BillItem, FeeItem, Product, Quote, Settings, Shipment, _qty, calc, check_bill, compare,
                             guess_category, parse_rule, weights)
from hs_agent.quotes import extract, parse_bill_lines, read_file, template_excel
from hs_agent.us_hts import UsHts, build, ch99_extra, merge_chapters, parse_rate, read_usitc

from pathlib import Path

DATA = Path(__file__).parent.parent / "hs_agent" / "data"

# USITC exportList 的 JSON 格式（字段名按官网导出）
RAW = [
    {"htsno": "8516", "indent": "0", "description": "Electric instantaneous or storage water heaters ...", "general": "", "footnotes": []},
    {"htsno": "8516.60", "indent": "1", "description": "Other ovens; cooking stoves, ranges, ...", "general": ""},
    {"htsno": "8516.60.40", "indent": "2", "description": "Cooking stoves, ranges and ovens", "general": "Free",
     "units": [], "footnotes": [{"columns": ["general"], "value": "See 9903.88.03.", "type": "endnote"}]},
    {"htsno": "8516.60.40.40", "indent": "3", "description": "Ranges", "general": "", "units": ["No."]},
    {"htsno": "8516.60.40.60", "indent": "3", "description": "Other", "general": "", "units": ["No."]},
    {"htsno": "", "indent": "2", "description": "Other:", "general": ""},
    {"htsno": "8516.60.60", "indent": "3", "description": "Roasters", "general": "4%", "units": ["No."],
     "footnotes": [{"columns": ["general"], "value": "See 9903.88.15 and 9903.01.25.", "type": "endnote"}]},
    {"htsno": "9403", "indent": "0", "description": "Other furniture and parts thereof:", "general": ""},
    {"htsno": "9403.20.00", "indent": "1", "description": "Other metal furniture", "general": "Free",
     "footnotes": [{"columns": ["general"], "value": "See 9903.88.03.", "type": "endnote"}]},
    {"htsno": "9903.88.03", "indent": "1", "description": "Articles the product of China, as provided for in U.S. note 20(e)",
     "general": "The duty provided in the applicable subheading + 25%"},
    {"htsno": "9903.88.15", "indent": "1", "description": "List 4A", "general": "The duty provided in the applicable subheading + 7.5%"},
    {"htsno": "9903.05.31", "indent": "1", "description": "Products of China", "general": "The duty provided in the applicable subheading + 12.5%"},
]


@pytest.fixture
def hts():
    return build(RAW, "test")


@pytest.fixture
def policy(tmp_path):
    ex = tmp_path / "ex.csv"
    ex.write_text("code,note\n85166060,测试豁免\n", encoding="utf-8")
    return Policy(DATA / "us_policy.json", ex)


def test_parse_rate():
    assert parse_rate("Free").ad_valorem == 0 and parse_rate("Free").ok
    assert parse_rate("2.7%").ad_valorem == 2.7
    r = parse_rate("4.4¢/kg + 5%")
    assert r.ad_valorem == 5 and r.specific == [(0.044, "kg")] and r.ok
    assert parse_rate("$1.25/doz").specific == [(1.25, "doz")]
    assert not parse_rate("The rate applicable to the natural juice in heading 2009").ok
    assert ch99_extra("The duty provided in the applicable subheading + 25%") == 25
    assert ch99_extra("The duty provided in the applicable subheading") == 0
    assert ch99_extra("Free") is None


def test_hts_lookup_inherits_rate_and_footnotes(hts):
    lines = hts.under("851660")
    assert [l.public()["code"] for l in lines] == ["8516.60.40.40", "8516.60.40.60", "8516.60.60"]
    other = lines[1].public()
    assert other["general"] == "Free" and other["rate_from"] == "8516.60.40"
    assert other["refs"][0]["ref"] == "9903.88.03" and other["refs"][0]["extra"] == 25
    roaster = lines[2].public()
    assert roaster["general"] == "4%" and "Other:" in roaster["full_desc"]
    assert {r["kind"] for r in roaster["refs"]} == {"301", "ieepa"}


def test_read_csv_and_merge(tmp_path):
    csv_text = "HTS Number,Indent,Description,Unit of Quantity,General Rate of Duty,Special Rate of Duty,Column 2 Rate of Duty\n" \
               "6109.10.00,1,Of cotton,,16.5%,Free (AU),90%\n"
    raw = read_usitc(csv_text.encode(), "a.csv")
    new = build(raw, "csv")
    assert new.get("61091000").general.ad_valorem == 16.5
    merged = merge_chapters(build(RAW, "x").lines, new.lines)
    assert [l.code[:2] for l in merged if l.code][0] == "61"
    hts = UsHts.from_lines(merged)
    hts.save(tmp_path / "h.json")
    assert UsHts(tmp_path / "h.json").get("8516604060")


def test_duty_layers(hts, policy):
    day = date(2026, 11, 20)
    r = line_duty(hts, policy, "8516604060", 10000, 500, 1000, day)
    assert [l.name for l in r.layers] == ["MFN 基础税率", "Section 301", "Section 301 新加征"]
    assert r.total == pytest.approx(10000 * 0.375)
    # 在豁免清单里：不计 12.5%；脚注里的 IEEPA 不计
    r = line_duty(hts, policy, "8516606000", 10000, 500, 1000, day)
    assert r.rate == pytest.approx(4 + 7.5) and any("豁免" in w for w in r.warnings) and any("IEEPA" in w for w in r.warnings)
    # 232：不叠加 12.5%
    r = line_duty(hts, policy, "94032000", 10000, 10, 100, day, s232_rate=25)
    assert r.rate == pytest.approx(25 + 25) and any("不叠加" in w for w in r.warnings)
    assert any("232" in w for w in line_duty(hts, policy, "94032000", 1, 1, 1, day).warnings)
    # 税则里没有的税号
    assert "没有" in line_duty(hts, policy, "12345678", 100, 1, 1, day).warnings[0]


def test_entry_fees(policy):
    layers, warn = entry_fees(policy, 5000, True, date(2026, 11, 1))
    assert layers[0].amount == pytest.approx(34.58) and layers[1].amount == pytest.approx(6.25) and not warn
    layers, _ = entry_fees(policy, 1_000_000, False, date(2026, 11, 1))
    assert len(layers) == 1 and layers[0].amount == pytest.approx(670.86)
    assert entry_fees(policy, 1000, True, date(2026, 11, 1))[1]


def shipment(**kw):
    p = dict(name="空气炸锅", cn_code="8516601000", us_hts="8516.60.40.60", units=400, unit_value_usd=15, unit_cost_cny=120,
             units_per_carton=4, carton_l=60, carton_w=40, carton_h=40, carton_kg=18)
    p.update(kw)
    return Shipment(products=[Product(**p)], settings=Settings(usd_cny=7, arrival=date(2026, 11, 20), annual_rate_pct=0))


def sea_quote(**kw):
    q = dict(id="a", forwarder="甲", transport="sea", clearance="self", transit_days=25, items=[
        FeeItem(name="海派运费", driver="per_kg", rate=9, rate_high=11),
        FeeItem(name="报关费", driver="per_shipment", rate=150),
        FeeItem(name="清关费", driver="per_shipment", rate=100, currency="USD"),
        FeeItem(name="ISF", driver="per_shipment", rate=35, currency="USD"),
        FeeItem(name="卡派", driver="per_cbm", rate=260, min_charge=800),
        FeeItem(name="USTR 港口费", driver="per_cbm", rate=15, scenario=True),
    ])
    q.update(kw)
    return Quote(**q)


def cn_table(tmp_path):
    f = tmp_path / "cn.csv"
    f.write_text("code,name,refund_rate,vat_rate,conditions,note,version,source\n8516601000,电烤箱,13,13,,,v,s\n", encoding="utf-8")
    return CnTariff(f)


def test_weights_and_fenpao():
    s = shipment()
    w = weights(s, Quote(divisor=6000))
    assert w["cartons"] == 100 and w["cbm"] == pytest.approx(9.6) and w["volume_kg"] == 1600 and w["chargeable_kg"] == 1800
    s = shipment(carton_kg=10)
    assert weights(s, Quote(divisor=6000))["chargeable_kg"] == 1600
    assert weights(s, Quote(divisor=6000, fenpao=0.5))["chargeable_kg"] == 1300  # 抛重 600 让一半
    assert weights(s, Quote(divisor=6000, min_kg=5000))["chargeable_kg"] == 5000


def test_carton_min_and_rules():
    # 两种箱子：轻小箱 50 箱（30×30×30cm、5kg），大重箱 10 箱（130×50×40cm、30kg）
    s = Shipment(products=[Product(units=50, unit_value_usd=1, carton_l=30, carton_w=30, carton_h=30, carton_kg=5),
                           Product(units=10, unit_value_usd=1, carton_l=130, carton_w=50, carton_h=40, carton_kg=30)],
                 settings=Settings(usd_cny=7))
    # 设了单箱最低就逐箱算：轻箱每箱 max(5, 4.5, 12)=12，重箱每箱 max(30, 43.33)=43.33
    assert weights(s, Quote(divisor=6000, min_kg_carton=12))["chargeable_kg"] == math.ceil(50 * 12 + 10 * 260000 / 6000)
    w = weights(s, Quote(divisor=6000))
    over = FeeItem(name="超重费", driver="per_carton", rate=180, carton_rule="重>22")
    long_ = FeeItem(name="超长费", driver="per_carton", rate=180, carton_rule="边>120 或 围>260")
    big = FeeItem(name="超大件", driver="per_carton", rate=1300, carton_rule="围>=320 且 围<400 或 边>=175 且 边<240")
    assert _qty(over, w, 0)[0] == 10 and _qty(long_, w, 0)[0] == 10 and _qty(big, w, 0)[0] == 0
    assert _qty(FeeItem(name="贴标", driver="per_carton", rate=1), w, 0)[0] == 60
    assert parse_rule("单箱实重>22kg 或 最长边≥120cm") == [[("kg", ">", 22)], [("side", ">=", 120)]]


def test_packing_checks():
    s = shipment(carton_kg=25, carton_l=70)
    texts = [c["text"] for c in calc(s, sea_quote(), build(RAW, "t"), Policy(DATA / "us_policy.json"), CnTariff(Path("none")))["checks"]]
    assert any("50 磅" in t for t in texts) and any("25 英寸" in t for t in texts)
    s = shipment(units_per_carton=1, units=10, carton_kg=30)
    texts = [c["text"] for c in calc(s, sea_quote(), build(RAW, "t"), Policy(DATA / "us_policy.json"), CnTariff(Path("none")))["checks"]]
    assert any("Team Lift" in t for t in texts)


def test_calc_self_clear(tmp_path, hts, policy):
    r = calc(shipment(), sea_quote(), hts, policy, cn_table(tmp_path))
    items = {i["name"]: i for i in r["items"]}
    assert items["海派运费"]["cny"] == 9 * 1800 and items["海派运费"]["cny_high"] == 11 * 1800
    assert items["清关费"]["cny"] == 700 and items["卡派"]["cny"] == pytest.approx(260 * 9.6)
    assert items["USTR 港口费"]["cny"] == 0 and items["USTR 港口费"]["cny_high"] == pytest.approx(144)
    duty = next(i for i in r["items"] if i["name"].startswith("关税"))
    assert duty["usd"] == pytest.approx(6000 * 0.375)
    assert {i["name"] for i in r["items"]} >= {"MPF 货物处理费", "HMF 港口维护费"}
    refund = r["refunds"][0]
    assert refund["rate"] == 13 and refund["net"] == pytest.approx(48000 / 1.13 * 0.13, abs=0.01)
    t = r["totals"]
    assert t["landed"] == pytest.approx(t["cost"] + 48000 - refund["net"], abs=0.05)
    pp = r["per_product"][0]
    assert pp["landed_per_unit"] == pytest.approx(t["landed"] / 400, abs=0.01)
    assert pp["landed_per_unit_high"] > pp["landed_per_unit"]


def test_calc_ddp(tmp_path, hts, policy):
    q = Quote(id="b", forwarder="乙", clearance="ddp", transit_days=35, includes=["清关费", "关税和海关规费", "派送费"],
              items=[FeeItem(name="包税海派", driver="per_kg", rate=14)])
    r = calc(shipment(), q, hts, policy, cn_table(tmp_path))
    assert not any(i["category"] == "关税和海关规费" for i in r["items"])
    assert r["refunds"][0]["net"] == 0 and "包税" in r["refunds"][0]["note"]
    assert r["implied_ddp"]["per_kg"] > 0


def test_compare_flags(tmp_path, hts, policy):
    s, cn = shipment(), cn_table(tmp_path)
    quotes = [
        sea_quote(),
        sea_quote(id="b", forwarder="乙", transit_days=30, items=[FeeItem(name="海派运费", driver="per_kg", rate=8.5),
                                                                FeeItem(name="清关费", driver="per_shipment", rate=100, currency="USD"),
                                                                FeeItem(name="卡派", driver="per_cbm", rate=600)]),
        sea_quote(id="c", forwarder="丙", transit_days=28, items=[FeeItem(name="海派运费", driver="per_kg", rate=9.2),
                                                                FeeItem(name="报关费", driver="per_shipment", rate=150),
                                                                FeeItem(name="清关费", driver="per_shipment", rate=90, currency="USD"),
                                                                FeeItem(name="ISF", driver="per_shipment", rate=35, currency="USD"),
                                                                FeeItem(name="卡派", driver="per_cbm", rate=250)]),
    ]
    rows = compare([calc(s, q, hts, policy, cn) for q in quotes], quotes, daily_profit=100)
    b = rows[1]
    texts = " ".join(f["text"] for f in b["flags"])
    assert "没写「报关费」" in texts and "没写「ISF/AMS」" in texts and "派送费" in texts and "高" in texts
    assert rows[0]["opportunity"] == 0 and b["opportunity"] == 500
    assert not any("没写" in f["text"] for f in rows[0]["flags"])


def test_compare_vague_fees(tmp_path, hts, policy):
    s, cn = shipment(), cn_table(tmp_path)
    q = sea_quote(note="目的港费用实报实销")
    q.items.append(FeeItem(name="查验费", note="按实际另收"))
    q.items.append(FeeItem(name="换单费", rate=300))
    r = calc(s, q, hts, policy, cn)
    texts = " ".join(f["text"] for f in compare([r], [q])[0]["flags"])
    assert "「查验费」没给价格" in texts and "实报实销" in texts
    assert {i["name"]: i["category"] for i in r["items"]}["换单费"] == "港口费"


def test_bill_check(tmp_path, hts, policy):
    r = calc(shipment(), sea_quote(), hts, policy, cn_table(tmp_path))
    bill = parse_bill_lines("海派运费 16200\n报关费 150\n清关费 USD 100\nISF 35 USD\n卡派 2496\n查验费 USD 300\n关税 USD 2250\nMPF 34.58 USD\nHMF 7.5 USD")
    assert bill and bill[2].currency == "USD" and bill[3].currency == "USD"
    out = check_bill(r, bill, 7)
    rows = {x["category"]: x for x in out["rows"]}
    assert rows["查验费"]["flag"] == "报价外费用" and rows["主干运费"]["flag"] == ""
    assert out["billed_total"] > out["quoted_total"]
    assert parse_bill_lines("这是一段说明文字") is None
    assert guess_category("旺季附加费") == "旺季附加" and guess_category("美国关税垫付手续费") == "关税垫付费"


def test_template_roundtrip():
    quotes, text = read_file(template_excel(), "报价模板.xlsx")
    assert text == "" and len(quotes) == 2
    a, b = quotes
    assert a.forwarder == "示例货代" and len(a.items) == 7 and a.items[2].currency == "USD" and a.items[4].min_charge == 800
    assert a.min_kg_carton == 12 and a.items[5].carton_rule == "重>22" and a.items[6].driver == "per_carton"
    assert a.items[0].rate_high == 11 and a.transit_days_high == 35
    assert b.clearance == "ddp" and "派送费" in b.includes and b.min_kg == 100
    ws = load_workbook(io.BytesIO(template_excel())).worksheets[1]
    assert "计费方式" in ws["A3"].value


def test_ai_extract_quote():
    reply = {"quotes": [{"forwarder": "丁", "channel": "美森快船", "transport": "sea", "clearance": "self", "transit_days": 18,
                         "transit_days_high": 25, "fenpao": 50, "includes": ["派送费", "瞎写的类别"],
                         "items": [{"name": "海派单价 100kg+", "category": "主干运费", "driver": "per_kg", "rate": "10.5", "currency": "CNY",
                                    "min_charge": 0, "rate_high": None, "note": "选了 100kg+ 档"},
                                   {"name": "查验费", "category": "乱写", "driver": "每票", "rate": 0, "note": "按实际另收"}]}]}
    seen = []

    async def complete(messages):
        seen.append(messages[0]["content"])
        return json.dumps(reply, ensure_ascii=False), 100, 50

    quotes, i, o = asyncio.run(extract("海派 10.5/kg ...", shipment(), complete))
    q = quotes[0]
    assert q.fenpao == 0.5 and q.includes == ["派送费"] and q.items[0].rate == 10.5
    assert q.items[1].category == "查验费" and q.items[1].driver == "per_shipment"
    assert "100 箱" in seen[0] and i == 100


def test_landed_over_http(tmp_path):
    import base64
    import threading
    import urllib.request

    from hs_agent.config import Config
    from hs_agent.web import App, serve

    port = 8797
    cfg = Config()
    cfg.notes_online = False
    cfg.data_dir = tmp_path
    cfg.us_hts_file = tmp_path / "us_hts.json"
    cfg.us_exempt_file = tmp_path / "ex.csv"
    reply = {"quotes": [{"forwarder": "丁", "transport": "sea", "clearance": "self", "transit_days": 20,
                         "items": [{"name": "海派", "category": "主干运费", "driver": "per_kg", "rate": 10, "currency": "CNY"}]}]}

    async def complete(messages):
        return json.dumps(reply, ensure_ascii=False), 10, 10

    app = App(cfg, completer=complete)
    srv = serve(app, port)
    threading.Thread(target=srv.serve_forever, daemon=True).start()

    def call(path, body=None):
        data = json.dumps(body).encode() if body is not None else None
        req = urllib.request.Request(f"http://127.0.0.1:{port}{path}", data=data, headers={"Content-Type": "application/json"})
        with urllib.request.urlopen(req, timeout=10) as r:
            raw = r.read()
            return raw if r.headers.get_content_type().endswith("sheet") else json.loads(raw)

    try:
        assert call("/api/landed/status")["us_hts"]["loaded"] is False
        build(RAW, "test").save(cfg.us_hts_file)  # 下载完不用重启
        assert call("/api/landed/status")["us_hts"]["count"] == len(RAW)
        look = call("/api/landed/lookup?code=8516601000")
        assert [u["code"] for u in look["us_lines"]][:1] == ["8516.60.40.40"]
        got = call("/api/quotes/extract", {"filename": "t.xlsx", "data": base64.b64encode(call("/api/quotes/template")).decode()})
        assert got["via"] == "template" and len(got["quotes"]) == 2
        got = call("/api/quotes/extract", {"text": "海派 10 元/kg", "shipment": shipment().model_dump(mode="json")})
        assert got["via"] == "ai" and got["quotes"][0]["forwarder"] == "丁"
        body = {"shipment": shipment().model_dump(mode="json"), "quotes": [sea_quote().model_dump(), got["quotes"][0]]}
        r = call("/api/landed/calc", body)
        assert len(r["results"]) == 2 and r["compare"][1]["flags"]
        r = call("/api/landed/bill", {"shipment": body["shipment"], "quote": body["quotes"][0], "text": "海派运费 16200\n查验费 USD 300"})
        assert r["record_id"] == 1 and any(x["flag"] == "报价外费用" for x in r["rows"])
        assert call("/api/landed/records")["count"] == 1
        with pytest.raises(urllib.error.HTTPError) as e:
            call("/api/landed/calc", {"shipment": {"products": [{"units": 0, "unit_value_usd": 1}]}, "quotes": [{}]})
        assert e.value.code == 400
    finally:
        srv.shutdown()


NOTES_TEXT = """20. (a)  For the purposes of heading 9903.88.03, products of China ...
    (f)    Heading 9903.88.03 applies to all products of China that are classified in the following 8-digit subheadings, except
           products that are provided for in headings 9903.88.33 ...
                 8516.60.40   9403.20.00   6307.90.9870
                        Harmonized Tariff Schedule of the United States Revision 21 (2026)
                 9401.71.00
    (g)    For the purposes of heading 9903.88.04, products of China ...
                 8516.79.00
52.
      (b)        As provided in heading 9903.05.86, the duties imposed by headings 9903.05.20–9903.05.84 shall not apply to articles
                 that are classifiable in the following provisions of the HTSUS:
                 0901.21.00   8471.30.01
    (c)         As provided in heading 9903.05.87, the duties imposed by headings 9903.05.20–9903.05.84 shall not apply to the following
                 0805.90.01
    (d)    As provided in heading 9903.05.88, the additional duties imposed by headings 9903.05.20–9903.05.84 shall not apply to
                 9403.20.00
"""


def test_parse_ch99_notes(tmp_path):
    from hs_agent.us_notes import parse, write
    s301, exempt = parse(NOTES_TEXT)
    assert [(r["code"], r["heading"]) for r in s301] == [("85166040", "9903.88.03"), ("94032000", "9903.88.03"),
                                                         ("6307909870", "9903.88.03"), ("94017100", "9903.88.03")]
    assert [(r["code"], r["conditional"]) for r in exempt] == [("09012100", 0), ("84713001", 0), ("94032000", 1)]
    write(s301, tmp_path / "s301.csv", ["code", "heading", "list"])
    write(exempt, tmp_path / "ex.csv", ["code", "heading", "note", "conditional"])
    pol = Policy(DATA / "us_policy.json", tmp_path / "ex.csv", tmp_path / "s301.csv")
    hts = build([r for r in RAW if r["htsno"] != "8516.60.40" or True], "t")
    for line in hts.lines:  # 去掉脚注，只靠清单判断 301
        line.footnotes = []
    day = date(2026, 11, 20)
    r = line_duty(hts, pol, "8516604060", 1000, 10, 10, day)
    assert [l.name for l in r.layers] == ["MFN 基础税率", "Section 301", "Section 301 新加征"]
    r = line_duty(hts, pol, "94032000", 1000, 10, 10, day)
    assert r.rate == 25 + 12.5 and any("有条件豁免" in w for w in r.warnings)
    assert ch99_extra("The duty provided in the applicable subheading plus 25%") == 25
