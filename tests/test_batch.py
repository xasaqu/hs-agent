"""批量归类：读清单、拼描述、对比已有编码、导出 Excel、经由网页上传和下载。"""

import asyncio
import base64
import csv
import io
import json
import threading
import time
import urllib.request
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from hs_agent.agent import RunResult
from hs_agent.batch import Batch, BatchError, compare, read_table, template_excel
from hs_agent.config import Config
from hs_agent.toolbox import Classification, Step
from hs_agent.web import App, serve

PORT = 8798


def xlsx(rows) -> bytes:
    wb = Workbook()
    for r in rows:
        wb.active.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def answer(hs6="851660", cn="8516601000", conf=0.85, review=False) -> Classification:
    return Classification(
        hs6=hs6, hs6_desc="Ovens", path="", cn_code=cn, cn_name="电烤箱", refund_rate=13.0, vat_rate="13",
        data_version="2026C", product_summary="空气炸锅",
        steps=[Step(code="8516", reason="家用电热器具"), Step(code="851660", reason="烹饪用")],
        rules=["规则一"], confidence=conf, need_review=review, missing_info=["功率"], alternatives=["851679"],
    )


def fake_classify(seen: list):
    async def fn(description, config, answerer=None, model=None, on_event=None):
        seen.append(description)
        await asyncio.sleep(0.01)
        if "坏" in description:
            return RunResult("x", Path("x.jsonl"), None, 3, 0, 0, 1.0, 0, 0, 0.01, error="模型超时")
        review = "T恤" in description
        a = answer("610910", "6109100021", 0.6, True) if review else answer()
        return RunResult("x", Path("x.jsonl"), a, 5, 0, 0, 1.0, 100, 50, 0.1)
    return fn


def cfg(tmp_path) -> Config:
    c = Config()
    c.notes_online = False
    c.data_dir = tmp_path
    return c


def test_read_table_columns():
    data = xlsx([["序号", "商品名称", "规格型号", "数量", "HS编码", "用途"],
                 [1, "空气炸锅", "5L", 10, "8516.60.1000", "家用"],
                 [2, "男式T恤", None, 5, None, None]])
    sheet = read_table(data, "清单.xlsx")
    assert [sheet.headers[i] for i in sheet.used_cols] == ["商品名称", "规格型号", "用途"]
    assert sheet.description(sheet.rows[0]) == "空气炸锅。规格型号：5L。用途：家用"
    assert sheet.description(sheet.rows[1]) == "男式T恤"
    assert sheet.given_code(sheet.rows[0]) == "8516601000" and sheet.given_code(sheet.rows[1]) == ""


def test_read_csv_gbk_and_errors():
    text = "商品名称,商品描述\n电动牙刷,声波震动\n"
    sheet = read_table(text.encode("gb18030"), "a.csv")
    assert sheet.description(sheet.rows[0]) == "电动牙刷。商品描述：声波震动"
    with pytest.raises(BatchError, match="xlsx"):
        read_table(b"x", "a.xls")
    with pytest.raises(BatchError, match="没有数据"):
        read_table(xlsx([["商品名称"]]), "a.xlsx")


def test_compare():
    assert compare("8516601000", "8516.60.1000".replace(".", "")) == "一致"
    assert compare("851660", "8516601000") == "一致"
    assert compare("8516609000", "8516601000") == "6 位一致，后几位不同"
    assert compare("8516790000", "8516601000") == "品目一致，子目不同"
    assert compare("9403", "8516601000") == "不一致"
    assert compare("", "8516601000") == ""


def test_batch_run_and_excel(tmp_path):
    data = xlsx([["商品名称", "规格", "HS编码"],
                 ["空气炸锅", "5L", "8516.60.9000"],
                 ["男式T恤", "棉", ""],
                 ["空气炸锅", "5L", ""],      # 和第 1 行相同，只归类一次
                 ["坏数据", "", ""],
                 ["", "", "1234"]])          # 没有商品信息
    seen = []
    b = Batch(read_table(data, "a.xlsx"), cfg(tmp_path), classify_fn=fake_classify(seen))
    asyncio.run(b.run())
    assert sorted(seen) == sorted(["空气炸锅。规格：5L", "男式T恤。规格：棉", "坏数据"])
    assert [r.status for r in b.rows] == ["done", "done", "done", "error", "skipped"]
    c = b.counts()
    assert c["done"] == 3 and c["review"] == 2 and c["cost_yuan"] == 0.21  # 出错的那行也有费用

    ws = load_workbook(io.BytesIO(b.to_excel())).worksheets[0]
    head = [c.value for c in ws[1]]
    assert head[:3] == ["商品名称", "规格", "HS编码"] and "与原编码对比" in head
    row = dict(zip(head, [c.value for c in ws[2]]))
    assert row["建议编码"] == "8516.60.1000" and row["退税率"] == "13%" and row["需人工复核"] == "否"
    assert row["与原编码对比"] == "6 位一致，后几位不同" and "8516：家用电热器具" in row["归类理由"]
    t = dict(zip(head, [c.value for c in ws[3]]))
    assert t["需人工复核"] == "是" and ws[3][3].fill.fgColor.rgb.endswith("FFF4D6")
    assert "相同" in dict(zip(head, [c.value for c in ws[4]]))["备注"]
    assert dict(zip(head, [c.value for c in ws[5]]))["备注"] == "模型超时"
    assert load_workbook(io.BytesIO(b.to_excel())).worksheets[1]["B1"].value == 5


def test_cancel_skips_queued_rows(tmp_path):
    data = xlsx([["商品名称"]] + [[f"商品{i}"] for i in range(6)])
    seen = []
    b = Batch(read_table(data, "a.xlsx"), cfg(tmp_path), classify_fn=fake_classify(seen), concurrency=1)
    b.cancel()
    asyncio.run(b.run())
    assert seen == [] and all(r.status == "skipped" for r in b.rows)


def test_template_is_readable():
    sheet = read_table(template_excel(), "模板.xlsx")
    assert sheet.code_col is not None and len(sheet.rows) == 2


def call(path, body=None):
    data = json.dumps(body).encode() if body is not None else None
    req = urllib.request.Request(f"http://127.0.0.1:{PORT}{path}", data=data, headers={"Content-Type": "application/json"})
    with urllib.request.urlopen(req, timeout=10) as r:
        raw = r.read()
        return raw if r.headers.get_content_type().endswith("sheet") else json.loads(raw)


def test_batch_over_http(tmp_path):
    seen = []
    app = App(cfg(tmp_path), classify_fn=fake_classify(seen))
    srv = serve(app, PORT)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    try:
        assert read_table(call("/api/batch/template"), "t.xlsx").rows
        data = xlsx([["商品名称", "用途"], ["空气炸锅", "家用"], ["男式T恤", "穿着"]])
        job = call("/api/batch", {"filename": "清单.xlsx", "data": base64.b64encode(data).decode()})
        assert job["counts"]["total"] == 2 and job["columns"] == ["商品名称", "用途"]
        for _ in range(100):
            job = call(f"/api/batch/{job['id']}")
            if job["status"] != "running":
                break
            time.sleep(0.05)
        assert job["status"] == "done" and [r["code"] for r in job["rows"]] == ["8516.60.1000", "6109.10.0021"]
        ws = load_workbook(io.BytesIO(call(f"/api/batch/{job['id']}/download"))).worksheets[0]
        assert ws.max_row == 3 and "与原编码对比" not in [c.value for c in ws[1]]
        assert list((tmp_path / "batch").glob("*归类结果.xlsx"))
        with pytest.raises(urllib.error.HTTPError) as e:
            call("/api/batch", {"filename": "a.xls", "data": base64.b64encode(b"x").decode()})
        assert e.value.code == 400
    finally:
        srv.shutdown()
