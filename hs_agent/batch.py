"""批量归类：读一张商品清单（Excel 或 CSV），逐行归类，导出一张带编码、退税率和依据的 Excel。

  python -m hs_agent.batch 商品清单.xlsx               # 结果写到 agent_data/batch/
  python -m hs_agent.batch 商品清单.xlsx -o 结果.xlsx --limit 5

清单第一行是表头。常见的列名会自动识别：品名/商品名称、规格型号、商品描述、材质、用途……
所有识别不出用途的列也会写进描述里；序号、数量、单价、金额这类和归类无关的列会跳过。
如果清单里已经有 HS 编码列，它不会给 Agent 看，只用来和结果对比（检查已有申报）。
批量时不追问，信息不足的按常见情况假设，并标记「需人工复核」。
"""

from __future__ import annotations

import argparse
import asyncio
import csv
import io
import re
import time
from dataclasses import dataclass, field
from pathlib import Path
from typing import Awaitable, Callable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .agent import RunResult, classify
from .config import Config
from .tariff import fmt_code
from .toolbox import no_more_info

MAX_ROWS = 500
NAME_RE = re.compile(r"品名|商品名称|货物名称|名称|产品名|中文名|^name|product", re.I)
CODE_RE = re.compile(r"hs|商品编码|税号|海关编码|税则号|编码", re.I)
SKIP_RE = re.compile(r"^(序号|编号|行号|no\.?|#)$|数量|单价|金额|总价|币制|币种|单位|箱数|件数|毛重|净重|price|qty|amount", re.I)
OUT_COLS = [
    "建议编码", "中文名称", "HS 6 位", "退税率", "增值税率", "监管条件", "置信度", "需人工复核",
    "与原编码对比", "归类理由", "依据", "缺少的信息", "其他可能编码", "备注", "记录编号",
]
YELLOW = PatternFill("solid", fgColor="FFF4D6")
RED = PatternFill("solid", fgColor="FDE2E0")
HEAD = PatternFill("solid", fgColor="E8EEFC")

ClassifyFn = Callable[..., Awaitable[RunResult]]


class BatchError(Exception):
    pass


class _Stopped(Exception):
    pass


@dataclass
class Sheet:
    headers: list[str]
    rows: list[list[str]]
    name_col: int | None
    code_col: int | None
    used_cols: list[int]

    def description(self, row: list[str]) -> str:
        parts = []
        if self.name_col is not None and row[self.name_col]:
            parts.append(row[self.name_col])
        for i in self.used_cols:
            if i != self.name_col and row[i]:
                parts.append(f"{self.headers[i]}：{row[i]}")
        return "。".join(parts)[:2000]

    def given_code(self, row: list[str]) -> str:
        if self.code_col is None:
            return ""
        return re.sub(r"\D", "", row[self.code_col])


def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def read_table(data: bytes, filename: str) -> Sheet:
    """读 Excel（.xlsx）或 CSV，返回表头、数据行，以及哪些列用来拼描述。"""
    suffix = Path(filename).suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        try:
            ws = load_workbook(io.BytesIO(data), read_only=True, data_only=True).worksheets[0]
        except Exception as e:
            raise BatchError(f"打不开这个 Excel 文件：{e}") from e
        raw = [[_cell(v) for v in r] for r in ws.iter_rows(values_only=True)]
    elif suffix in (".csv", ".txt"):
        text = next((data.decode(enc) for enc in ("utf-8-sig", "gb18030") if _decodes(data, enc)), None)
        if text is None:
            raise BatchError("CSV 文件的编码无法识别，请另存为 UTF-8 或直接用 Excel 格式。")
        raw = [[_cell(v) for v in r] for r in csv.reader(io.StringIO(text))]
    elif suffix == ".xls":
        raise BatchError("不支持老版 .xls 格式，请在 Excel 里另存为 .xlsx 再上传。")
    else:
        raise BatchError("只支持 .xlsx 或 .csv 文件。")

    raw = [r for r in raw if any(r)]
    if len(raw) < 2:
        raise BatchError("表格里没有数据：第一行应是表头，下面每行一个商品。")
    width = max(len(r) for r in raw)
    raw = [r + [""] * (width - len(r)) for r in raw]
    headers = [h or f"第{i + 1}列" for i, h in enumerate(raw[0])]
    rows = raw[1:]
    if len(rows) > MAX_ROWS:
        raise BatchError(f"一次最多 {MAX_ROWS} 行，这个文件有 {len(rows)} 行，请分成几个文件。")

    out = set(OUT_COLS)
    name_col = next((i for i, h in enumerate(headers) if NAME_RE.search(h) and not CODE_RE.search(h)), None)
    code_col = next((i for i, h in enumerate(headers) if CODE_RE.search(h) and h not in out), None)
    used = [i for i, h in enumerate(headers)
            if i != code_col and h not in out and not SKIP_RE.search(h) and any(r[i] for r in rows)]
    if not used:
        raise BatchError("没有找到可以用来归类的列，请至少提供「商品名称」或「商品描述」。")
    return Sheet(headers, rows, name_col, code_col, used)


def _decodes(data: bytes, enc: str) -> bool:
    try:
        data.decode(enc)
        return True
    except UnicodeDecodeError:
        return False


def compare(given: str, got: str) -> str:
    if not given:
        return ""
    if not got:
        return "无法对比"
    n = min(len(given), len(got))
    if given[:n] == got[:n]:
        return "一致"
    if given[:6] == got[:6]:
        return "6 位一致，后几位不同"
    if given[:4] == got[:4]:
        return "品目一致，子目不同"
    return "不一致"


@dataclass
class RowResult:
    index: int  # 清单里的第几行（从 1 开始，不含表头）
    description: str
    status: str = "waiting"  # waiting / running / done / error / skipped
    values: dict = field(default_factory=dict)
    cost_yuan: float = 0.0
    log: str = ""


def row_values(r: RunResult, given: str) -> dict:
    a = r.answer
    if not a:
        return {"备注": r.error or "Agent 没有给出结果", "需人工复核": "是", "与原编码对比": compare(given, "")}
    code = a.cn_code or a.hs6
    rate = "" if a.refund_rate is None else f"{a.refund_rate:g}%"
    vat = a.vat_rate + ("%" if a.vat_rate and "/" not in a.vat_rate and not a.vat_rate.endswith("%") else "")
    notes = [a.refund_note] if a.refund_note else []
    if r.reused:
        notes.append(f"沿用{r.reused}，没有重新归类")
    if not a.cn_code:
        notes.append("未找到中国 10 位编码，只给出 6 位")
    return {
        "建议编码": fmt_code(code),
        "中文名称": a.cn_name,
        "HS 6 位": fmt_code(a.hs6),
        "退税率": rate if a.cn_code else "",
        "增值税率": vat,
        "监管条件": a.conditions,
        "置信度": round(a.confidence, 2),
        "需人工复核": "是" if a.need_review else "否",
        "与原编码对比": compare(given, code),
        "归类理由": "\n".join(f"{fmt_code(s.code)}：{s.reason}" for s in a.steps),
        "依据": "\n".join(a.rules),
        "缺少的信息": "\n".join(a.missing_info),
        "其他可能编码": "、".join(fmt_code(c) for c in a.alternatives),
        "备注": "；".join(notes),
        "记录编号": r.record_id or "",
        "_version": a.data_version,
    }


class Batch:
    """一次批量归类。可以在后台线程里跑，界面轮询 rows 看进度；cancel() 停止剩下的行。"""

    def __init__(self, sheet: Sheet, config: Config, classify_fn: ClassifyFn = classify, model=None,
                 concurrency: int = 2, limit: int = 0):
        self.sheet, self.config, self.classify_fn, self.model = sheet, config, classify_fn, model
        self.concurrency = max(1, concurrency)
        rows = sheet.rows[:limit] if limit else sheet.rows
        self.rows = [RowResult(i + 1, sheet.description(r)) for i, r in enumerate(rows)]
        self.cancelled = False
        self.started = time.time()
        self.finished = 0.0

    def cancel(self) -> None:
        self.cancelled = True

    async def run(self, on_row: Callable[[RowResult], None] | None = None) -> None:
        sem = asyncio.Semaphore(self.concurrency)
        same: dict[str, asyncio.Future] = {}  # 描述完全相同的行只归类一次

        async def classify_once(desc: str) -> RunResult:
            async with sem:
                if self.cancelled:  # 停止后，还在排队的行不再归类
                    raise _Stopped
                for r in self.rows:
                    if r.description == desc and r.status == "waiting":
                        r.status = "running"
                return await self.classify_fn(desc, self.config, answerer=no_more_info, model=self.model)

        async def one(row: RowResult) -> None:
            raw = self.sheet.rows[row.index - 1]
            if not row.description:
                row.status, row.values = "skipped", {"备注": "这一行没有可用的商品信息"}
            else:
                first = row.description not in same
                if first:
                    same[row.description] = asyncio.ensure_future(classify_once(row.description))
                try:
                    r = await same[row.description]
                    row.values = row_values(r, self.sheet.given_code(raw))
                    row.status = "done" if r.answer else "error"
                    row.cost_yuan = r.cost_yuan if first else 0.0
                    row.log = str(r.log_path)
                    if not first:
                        row.values["备注"] = "；".join(x for x in (row.values.get("备注", ""), "和前面描述相同的行结果一致") if x)
                except _Stopped:
                    row.status, row.values = "skipped", {"备注": "已停止，未归类"}
                except Exception as e:  # 一行出错不影响其他行
                    row.status, row.values = "error", {"备注": f"出错：{e}", "需人工复核": "是"}
            if on_row:
                on_row(row)

        await asyncio.gather(*(one(r) for r in self.rows))
        self.finished = time.time()

    # ---------- 汇总和导出 ----------

    def counts(self) -> dict:
        st = [r.status for r in self.rows]
        return {
            "total": len(self.rows), "done": st.count("done"), "error": st.count("error"),
            "skipped": st.count("skipped"), "running": st.count("running"),
            "review": sum(1 for r in self.rows if r.values.get("需人工复核") == "是"),
            "cost_yuan": round(sum(r.cost_yuan for r in self.rows), 4),
        }

    def to_excel(self) -> bytes:
        wb = Workbook()
        ws = wb.active
        ws.title = "归类结果"
        has_code = self.sheet.code_col is not None
        cols = [c for c in OUT_COLS if has_code or c != "与原编码对比"]
        n_in = len(self.sheet.headers)
        ws.append(self.sheet.headers + cols)
        for i, c in enumerate(ws[1], start=1):
            c.font = Font(bold=True)
            c.alignment = Alignment(vertical="center", wrap_text=True)
            if i > n_in:
                c.fill = HEAD
        for row in self.rows:
            raw = self.sheet.rows[row.index - 1]
            ws.append(raw + [row.values.get(c, "") for c in cols])
            fill = RED if row.status in ("error", "skipped") else YELLOW if row.values.get("需人工复核") == "是" else None
            for j, cell in enumerate(ws[ws.max_row], start=1):
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if fill and j > n_in:
                    cell.fill = fill
        widths = {"建议编码": 16, "HS 6 位": 10, "中文名称": 22, "退税率": 8, "增值税率": 9, "监管条件": 10, "置信度": 8,
                  "需人工复核": 10, "与原编码对比": 14, "归类理由": 60, "依据": 40, "缺少的信息": 40, "其他可能编码": 18, "备注": 24}
        for j, h in enumerate(self.sheet.headers + cols, start=1):
            w = widths.get(h) if j > n_in else min(40, max(8, max((len(str(r[j - 1])) for r in self.sheet.rows[:50]), default=8)))
            ws.column_dimensions[get_column_letter(j)].width = w or 12
        ws.freeze_panes = ws.cell(row=2, column=(self.sheet.name_col or 0) + 2)

        c = self.counts()
        versions = sorted({r.values.get("_version") for r in self.rows if r.values.get("_version")})
        info = wb.create_sheet("说明")
        for line in [
            ("商品数", c["total"]), ("完成归类", c["done"]), ("需人工复核", c["review"]),
            ("未完成（出错或跳过）", c["error"] + c["skipped"]), ("大模型费用（元，约）", c["cost_yuan"]),
            ("退税率数据版本", "、".join(versions) or "未导入"), ("模型", self.config.llm_model),
            ("生成时间", time.strftime("%Y-%m-%d %H:%M")),
            ("", ""),
            ("颜色", "黄色 = 建议人工复核；红色 = 没有完成归类"),
            ("用来归类的列", "、".join(self.sheet.headers[i] for i in self.sheet.used_cols)),
            ("提示", "批量时不追问，信息不足的按常见情况假设，写在「缺少的信息」里。补充信息后可以单独在网页上重新归类。"),
            ("声明", "结果仅供参考，不构成归类依据，正式结论以海关审核或预裁定为准。"),
        ]:
            info.append(line)
        info.column_dimensions["A"].width = 22
        info.column_dimensions["B"].width = 80
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()


def template_excel() -> bytes:
    """给用户下载的清单模板。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "商品清单"
    ws.append(["序号", "商品名称", "规格型号", "材质", "用途", "商品描述", "HS编码（可不填）"])
    ws.append([1, "空气炸锅", "5L，1500W", "塑料外壳", "家用烹饪", "电热管加热，风扇循环热风", ""])
    ws.append([2, "男式短袖T恤", "圆领", "100%棉，针织", "穿着", "", "6109100021"])
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = HEAD
    for col, w in zip("ABCDEFG", (6, 18, 14, 14, 12, 36, 18)):
        ws.column_dimensions[col].width = w
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def main() -> None:
    p = argparse.ArgumentParser(description="批量归类：Excel/CSV 商品清单 → 带编码和退税率的 Excel")
    p.add_argument("file")
    p.add_argument("-o", "--out", default="")
    p.add_argument("--limit", type=int, default=0, help="只跑前几行（先试一下）")
    p.add_argument("--concurrency", type=int, default=2)
    args = p.parse_args()
    config = Config()
    if not config.llm_api_key:
        raise SystemExit("缺少大模型 API Key：请在 .env 里设置 LLM_API_KEY（或 DEEPSEEK_API_KEY）。")
    src = Path(args.file)
    try:
        sheet = read_table(src.read_bytes(), src.name)
    except BatchError as e:
        raise SystemExit(str(e))
    batch = Batch(sheet, config, concurrency=args.concurrency, limit=args.limit)
    print(f"共 {len(batch.rows)} 行，用来归类的列：{'、'.join(sheet.headers[i] for i in sheet.used_cols)}")

    def show(row: RowResult) -> None:
        code = row.values.get("建议编码") or "—"
        flag = "（需复核）" if row.values.get("需人工复核") == "是" else ""
        print(f"[{sum(r.status in ('done', 'error', 'skipped') for r in batch.rows)}/{len(batch.rows)}] "
              f"第 {row.index} 行 {code}{flag} {row.values.get('备注', '')}", flush=True)

    asyncio.run(batch.run(show))
    out = Path(args.out) if args.out else config.data_dir / "batch" / f"{time.strftime('%Y%m%d-%H%M%S')}-{src.stem}-归类结果.xlsx"
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_bytes(batch.to_excel())
    c = batch.counts()
    print(f"\n完成 {c['done']}/{c['total']}，需人工复核 {c['review']}，费用约 {c['cost_yuan']} 元\n结果：{out}")


if __name__ == "__main__":
    main()
