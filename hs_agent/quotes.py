"""货代报价导入：按模板填的 Excel 直接读；其他格式的报价单（Excel、文字、聊天记录）交给大模型拆成统一的费用项。

另外保存「报价 + 实际账单」的对照记录（agent_data/logistics.db），攒够 30~50 票后用来看报价和实付差在哪里。
"""

from __future__ import annotations

import io
import json
import re
import sqlite3
import threading
import time
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from pydantic import ValidationError

from .baselines import Completer, _parse
from .landed import CATEGORY_NAMES, DRIVER_NAMES, BillItem, FeeItem, Quote, Shipment, guess_category, weights

MAX_TEXT = 40000

TEMPLATE_HEAD = ["货代", "渠道名称", "运输方式", "清关方式", "时效（天）", "最慢时效（天）", "体积重除数", "分泡让出比例",
                 "最低计费重（kg）", "单箱最低计费重（kg）", "计费重进位（kg）", "运费已含", "费用项", "费用类别", "计费方式",
                 "单价", "币种", "最低收费", "偏坏情况单价", "每箱条件", "备注"]
DRIVER_BY_NAME = {"每公斤": "per_kg", "每kg": "per_kg", "公斤": "per_kg", "kg": "per_kg", "每方": "per_cbm", "方": "per_cbm",
                  "cbm": "per_cbm", "每箱": "per_carton", "箱": "per_carton", "每票": "per_shipment", "票": "per_shipment",
                  "货值百分比": "pct_value", "运费百分比": "pct_freight", "百分比": "pct_value", "%": "pct_value", "每天": "per_day", "天": "per_day",
                  "每件": "per_unit", "件": "per_unit"}
DRIVER_BY_NAME.update({v.split("（")[0]: k for k, v in DRIVER_NAMES.items()})


class QuoteError(Exception):
    pass


def _driver(text: str) -> str:
    t = str(text or "").strip().lower()
    if t in DRIVER_NAMES:
        return t
    for k, v in DRIVER_BY_NAME.items():
        if k.lower() in t:
            return v
    return "per_shipment"


def _num(v, default=0.0) -> float:
    if v is None or str(v).strip() == "":
        return default
    m = re.search(r"-?\d+(?:\.\d+)?", str(v).replace(",", ""))
    return float(m.group(0)) if m else default


def template_excel() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "报价"
    ws.append(TEMPLATE_HEAD)
    rows = [
        ["示例货代", "美森海派自清", "海运", "自主清关", 25, 35, 6000, 0, 21, 12, 1, "", "海派运费（美西）", "主干运费", "每公斤", 9.5, "CNY", 0, 11, "", "21kg 以上档"],
        ["", "", "", "", "", "", "", "", "", "", "", "", "报关费", "报关费", "每票", 150, "CNY", 0, "", "", ""],
        ["", "", "", "", "", "", "", "", "", "", "", "", "清关费", "清关费", "每票", 100, "USD", 0, "", "", ""],
        ["", "", "", "", "", "", "", "", "", "", "", "", "ISF 申报", "ISF/AMS", "每票", 35, "USD", 0, "", "", ""],
        ["", "", "", "", "", "", "", "", "", "", "", "", "卡派到仓", "派送费", "每方", 260, "CNY", 800, 320, "", ""],
        ["", "", "", "", "", "", "", "", "", "", "", "", "超重费", "操作费", "每箱", 180, "CNY", 0, "", "重>22", "单箱实重超过 22kg"],
        ["", "", "", "", "", "", "", "", "", "", "", "", "超长费", "操作费", "每箱", 180, "CNY", 0, "", "边>120 或 围>260", "围长 = 最长边 + 2×(另两边)"],
        ["示例货代B", "普船海派包税", "海运", "包税", 35, 45, 6000, 0, 100, "", 1, "清关费、关税和海关规费、派送费", "包税海派单价", "主干运费", "每公斤", 14, "CNY", 0, "", "", "含税含派"],
    ]
    for r in rows:
        ws.append(r)
    head = PatternFill("solid", fgColor="E8EEF7")
    for c in ws[1]:
        c.font, c.fill = Font(bold=True), head
    for i, w in enumerate([12, 16, 8, 10, 9, 11, 10, 11, 12, 14, 12, 24, 18, 12, 12, 8, 6, 9, 12, 18, 22], 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    note = wb.create_sheet("说明")
    for line in [
        "一家货代的一个渠道占若干行：第一行填渠道信息（货代到运费已含），后面每行一个费用项，渠道信息可以留空。",
        "运输方式：海运 / 空运；清关方式：自主清关 / 包税。",
        "计费方式：" + "、".join(v.split("（")[0] for v in DRIVER_NAMES.values()) + "。货值百分比、运费百分比的单价填百分数，例如 0.3、25。燃油附加、旺季附加一般是运费百分比。",
        "费用类别：" + "、".join(CATEGORY_NAMES) + "。不填会按费用项名称自动判断。",
        "分泡让出比例：抛重部分让给客户的比例，例如五五分泡填 50%（或 0.5），不分泡填 0。",
        "运费已含：主运费里已经包含的费用类别，用顿号隔开，例如「清关费、派送费」。",
        "偏坏情况单价：旺季、涨价后的单价，不填则和单价相同。",
        "单箱最低计费重：例如「单箱不足 12kg 按 12kg」填 12；填了就逐箱算计费重再加总。",
        "每箱条件：超重费、超长费这类只对部分箱子收的每箱费用，写收费条件。重 = 单箱实重（kg），边 = 最长边（cm），"
        "围 = 最长边 + 2 ×（另两边之和）（cm）；「或」任一满足，「且」同时满足。例如：重>22；边>120 或 围>260。",
    ]:
        note.append([line])
    note.column_dimensions["A"].width = 120
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _sheets(data: bytes) -> list[list[list]]:
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as e:
        raise QuoteError(f"打不开这个 Excel 文件：{e}") from e
    return [[list(r) for r in ws.iter_rows(values_only=True)] for ws in wb.worksheets]


def _is_template(rows: list[list]) -> bool:
    head = [str(c or "").strip() for c in (rows[0] if rows else [])]
    return "费用项" in head and "计费方式" in head and "单价" in head


def read_template(rows: list[list]) -> list[Quote]:
    head = [str(c or "").strip() for c in rows[0]]
    col = {h: i for i, h in enumerate(head) if h}
    get = lambda r, h: r[col[h]] if h in col and col[h] < len(r) else None
    quotes: list[Quote] = []
    for r in rows[1:]:
        if not any(v not in (None, "") for v in r):
            continue
        if get(r, "货代") or get(r, "渠道名称") or not quotes:
            fenpao = _num(get(r, "分泡让出比例"))
            includes = [x for x in re.split(r"[、,，;；/\s]+", str(get(r, "运费已含") or "")) if x]
            quotes.append(Quote(
                id=f"q{len(quotes) + 1}", forwarder=str(get(r, "货代") or "").strip(), channel=str(get(r, "渠道名称") or "").strip(),
                transport="air" if "空" in str(get(r, "运输方式") or "") else "sea",
                clearance="ddp" if "包" in str(get(r, "清关方式") or "") or "双清" in str(get(r, "清关方式") or "") else "self",
                transit_days=_num(get(r, "时效（天）")), transit_days_high=_num(get(r, "最慢时效（天）")),
                divisor=_num(get(r, "体积重除数"), 6000) or 6000, fenpao=fenpao / 100 if fenpao > 1 else fenpao,
                min_kg=_num(get(r, "最低计费重（kg）")), min_kg_carton=_num(get(r, "单箱最低计费重（kg）")),
                round_kg=_num(get(r, "计费重进位（kg）"), 1),
                includes=[guess_category(x) if x not in CATEGORY_NAMES else x for x in includes],
            ))
        name = str(get(r, "费用项") or "").strip()
        if not name:
            continue
        high = get(r, "偏坏情况单价")
        cat = str(get(r, "费用类别") or "").strip()
        quotes[-1].items.append(FeeItem(
            name=name, category=cat if cat in CATEGORY_NAMES else guess_category(name), driver=_driver(get(r, "计费方式")),
            rate=_num(get(r, "单价")), currency="USD" if "USD" in str(get(r, "币种") or "").upper() or "美" in str(get(r, "币种") or "") else "CNY",
            min_charge=_num(get(r, "最低收费")), rate_high=_num(high) if high not in (None, "") else None,
            carton_rule=str(get(r, "每箱条件") or "").strip(), note=str(get(r, "备注") or "").strip(),
        ))
    return [q for q in quotes if q.items]


def rows_text(sheets: list[list[list]]) -> str:
    out = []
    for k, rows in enumerate(sheets, 1):
        out.append(f"--- 工作表 {k} ---")
        for r in rows:
            cells = [str(c).strip() for c in r if c not in (None, "")]
            if cells:
                out.append("\t".join(cells))
    return "\n".join(out)[:MAX_TEXT]


EXTRACT_PROMPT = """你是货代报价分析助手。把下面的头程报价（可能是表格、文字或聊天记录）拆成统一格式的 JSON。

这票货的情况（用来选价格档位）：{context}

输出 JSON：{{"quotes": [报价, ...]}}。一家货代的一个渠道是一个报价，字段：
- forwarder 货代名称；channel 渠道名称；transport "sea"（海运、海派、海卡）或 "air"（空运、空派）
- clearance "self"（自主清关、卖家做进口商、不包税）或 "ddp"（包税、双清包税、DDP）
- transit_days 常态时效天数，transit_days_high 最慢时效天数（写「25-35 天」就是 25 和 35）
- divisor 体积重除数（默认 6000）；fenpao 分泡时抛重让给客户的比例（五五分泡 0.5，不分泡 0）
- min_kg 整票最低计费重；min_kg_carton 单箱最低计费重（「单箱不足 12kg 按 12kg」就是 12，没有填 0）；round_kg 计费重进位（不写就 1）
- includes 主运费已经包含的费用类别（例如报价写「含税含派」就是 ["清关费", "关税和海关规费", "派送费"]）
- items 费用项列表，每项：name 原文名称，category 类别（只能从 {categories} 里选），
  driver 计费方式（只能是 per_kg 每公斤计费重、per_cbm 每方、per_carton 每箱、per_shipment 每票、pct_value 货值百分比（保险等）、pct_freight 运费百分比（燃油附加、旺季附加等按运费加收的）、per_day 每天、per_unit 每件），
  rate 单价（数字；货值百分比填百分数，如 0.3），currency "CNY" 或 "USD"，min_charge 最低收费（没有填 0），
  rate_high 旺季或上涨后的单价（报价里没有就填 null），note 备注（选了哪一档、有什么条件），
  carton_rule 只对满足条件的箱子收的每箱费用（超重、超长、超大件等）写条件，否则填 ""。
  条件写法：重（单箱实重 kg）、边（最长边 cm）、围（围长 = 最长边 + 2×另两边之和，cm），用 > >= < <= 比较；
  「或」连接任一满足即可，「且」连接要同时满足。例如「重>22」「边>120 或 围>260」「围>=320 且 围<400 或 边>=175 且 边<240」
- note 整个报价的备注：有效期、不包含的项目、特别条件；有免费仓储期 / 免费用箱期的，写明多少天、从哪天起算，以及超期怎么收

规则：
1. 价格按重量段、目的仓分档的，只取适合这票货的那一档，在 note 里写明选的哪一档；目的仓不明确时取美西。
2. 报价里只写了「另收」「按实际」而没有价格的费用，也要列出来，rate 填 0，note 写「按实际另收」。
3. 不要编造报价里没有的费用和价格。
4. 只输出 JSON。

报价原文：
{text}"""


def context_text(shipment: Shipment | None) -> str:
    if not shipment or not shipment.products:
        return "未提供"
    w = weights(shipment, Quote())
    dest = shipment.settings.destination or "未指定"
    return (f"{w['cartons']} 箱，实重 {w['actual_kg']}kg，体积 {w['cbm']} 方（按 /6000 体积重 {w['volume_kg']}kg），"
            f"货值 ${w['value_usd']:,.0f}，目的仓 {dest}")


def coerce_quote(d: dict, idx: int) -> Quote:
    items = []
    for it in d.get("items") or []:
        if not isinstance(it, dict) or not it.get("name"):
            continue
        cat = it.get("category") or ""
        items.append(FeeItem(
            name=str(it["name"])[:80], category=cat if cat in CATEGORY_NAMES else guess_category(str(it["name"])),
            driver=_driver(it.get("driver")), rate=_num(it.get("rate")),
            currency="USD" if str(it.get("currency", "")).upper() == "USD" else "CNY",
            min_charge=_num(it.get("min_charge")),
            rate_high=_num(it["rate_high"]) if it.get("rate_high") not in (None, "") else None,
            carton_rule=str(it.get("carton_rule") or "")[:80], note=str(it.get("note") or "")[:200],
        ))
    fenpao = _num(d.get("fenpao"))
    try:
        return Quote(
            id=f"q{idx}", forwarder=str(d.get("forwarder") or "")[:60], channel=str(d.get("channel") or "")[:60],
            transport="air" if d.get("transport") == "air" else "sea", clearance="ddp" if d.get("clearance") == "ddp" else "self",
            transit_days=_num(d.get("transit_days")), transit_days_high=_num(d.get("transit_days_high")),
            divisor=_num(d.get("divisor"), 6000) or 6000, fenpao=fenpao / 100 if fenpao > 1 else fenpao,
            min_kg=_num(d.get("min_kg")), min_kg_carton=_num(d.get("min_kg_carton")), round_kg=_num(d.get("round_kg"), 1) or 1,
            includes=[c for c in d.get("includes") or [] if c in CATEGORY_NAMES], items=items,
            note=str(d.get("note") or "")[:500],
        )
    except ValidationError as e:
        raise QuoteError(f"报价格式不对：{e}") from e


async def extract(text: str, shipment: Shipment | None, complete: Completer) -> tuple[list[Quote], int, int]:
    text = text.strip()[:MAX_TEXT]
    if not text:
        raise QuoteError("报价内容是空的。")
    prompt = EXTRACT_PROMPT.format(context=context_text(shipment), categories="、".join(CATEGORY_NAMES), text=text)
    reply, i, o = await complete([{"role": "user", "content": prompt}])
    try:
        data = _parse(reply)
    except json.JSONDecodeError as e:
        raise QuoteError("大模型返回的内容不是有效的 JSON，请重试。") from e
    quotes = [coerce_quote(d, k) for k, d in enumerate(data.get("quotes") or [], 1) if isinstance(d, dict)]
    quotes = [q for q in quotes if q.items]
    if not quotes:
        raise QuoteError("没能从这份内容里识别出报价。")
    return quotes, i, o


def read_file(data: bytes, filename: str) -> tuple[list[Quote] | None, str]:
    """模板格式直接返回报价；其他格式返回 (None, 文字内容)，交给大模型提取。"""
    name = filename.lower()
    if name.endswith((".xlsx", ".xlsm")):
        sheets = _sheets(data)
        if sheets and _is_template(sheets[0]):
            quotes = read_template(sheets[0])
            if not quotes:
                raise QuoteError("模板里没有读到费用项。")
            return quotes, ""
        return None, rows_text(sheets)
    if name.endswith((".csv", ".txt", ".md")):
        for enc in ("utf-8-sig", "gb18030"):
            try:
                return None, data.decode(enc)[:MAX_TEXT]
            except UnicodeDecodeError:
                continue
        raise QuoteError("文件编码无法识别。")
    if name.endswith(".xls"):
        raise QuoteError("不支持老版 .xls 格式，请在 Excel 里另存为 .xlsx。")
    raise QuoteError("报价文件只支持 .xlsx、.csv、.txt；其他格式可以直接把内容粘贴进来。")


BILL_PROMPT = """把下面的货代账单拆成费用明细，输出 JSON：{{"items": [{{"name": 原文费用名称, "amount": 金额数字, "currency": "CNY" 或 "USD", "category": 类别}}]}}。
类别只能从 {categories} 里选。只列实际收费的行，不要列合计行；不要编造。只输出 JSON。

账单原文：
{text}"""


async def extract_bill(text: str, complete: Completer) -> tuple[list[BillItem], int, int]:
    reply, i, o = await complete([{"role": "user", "content": BILL_PROMPT.format(categories="、".join(CATEGORY_NAMES), text=text[:MAX_TEXT])}])
    items = []
    for d in _parse(reply).get("items") or []:
        if isinstance(d, dict) and d.get("name") and _num(d.get("amount")):
            cat = d.get("category") or ""
            items.append(BillItem(name=str(d["name"])[:80], amount=_num(d["amount"]),
                                  currency="USD" if str(d.get("currency", "")).upper() == "USD" else "CNY",
                                  category=cat if cat in CATEGORY_NAMES else ""))
    if not items:
        raise QuoteError("没能从账单里识别出费用。")
    return items, i, o


def parse_bill_lines(text: str) -> list[BillItem] | None:
    """简单格式不用大模型：每行「费用名称 金额 [USD]」。有一行认不出就返回 None。"""
    items = []
    for line in text.strip().splitlines():
        line = line.strip()
        if not line:
            continue
        m = re.fullmatch(r"(.+?)[\s:：,，\t]+(?:(USD|\$|美元|CNY|RMB|￥|¥|元)\s*)?(-?[\d,]+(?:\.\d+)?)\s*(USD|美元|CNY|RMB|元)?", line, re.I)
        if not m:
            return None
        cur = (m.group(2) or m.group(4) or "").upper()
        items.append(BillItem(name=m.group(1).strip(), amount=float(m.group(3).replace(",", "")),
                              currency="USD" if cur in ("USD", "$", "美元") else "CNY"))
    return items or None


# ---------- 报价和账单对照记录 ----------

SCHEMA = """
CREATE TABLE IF NOT EXISTS shipments (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    created REAL NOT NULL,
    title TEXT NOT NULL,
    shipment TEXT NOT NULL,
    quote TEXT NOT NULL,
    result TEXT NOT NULL,
    bill TEXT NOT NULL,
    checked TEXT NOT NULL
);
"""


class Ledger:
    def __init__(self, path: Path):
        path.parent.mkdir(parents=True, exist_ok=True)
        self._lock = threading.Lock()
        self._db = sqlite3.connect(path, check_same_thread=False)
        self._db.row_factory = sqlite3.Row
        self._db.executescript(SCHEMA)

    def add(self, shipment: Shipment, quote: Quote, result: dict, bill: list[BillItem], checked: dict) -> int:
        with self._lock, self._db:
            cur = self._db.execute(
                "INSERT INTO shipments(created, title, shipment, quote, result, bill, checked) VALUES (?,?,?,?,?,?,?)",
                (time.time(), quote.title, shipment.model_dump_json(), quote.model_dump_json(), json.dumps(result, ensure_ascii=False),
                 json.dumps([b.model_dump() for b in bill], ensure_ascii=False), json.dumps(checked, ensure_ascii=False)))
            return cur.lastrowid

    def list(self, limit: int = 100) -> list[dict]:
        with self._lock:
            rows = self._db.execute("SELECT id, created, title, checked FROM shipments ORDER BY id DESC LIMIT ?", (limit,)).fetchall()
        out = []
        for r in rows:
            c = json.loads(r["checked"])
            out.append({"id": r["id"], "created": time.strftime("%Y-%m-%d %H:%M", time.localtime(r["created"])), "title": r["title"],
                        "quoted": c.get("quoted_total"), "billed": c.get("billed_total"), "diff_pct": c.get("diff_pct"),
                        "flags": [x["category"] + "：" + x["flag"] for x in c.get("rows", []) if x["flag"]]})
        return out

    def count(self) -> int:
        with self._lock:
            return self._db.execute("SELECT COUNT(*) FROM shipments").fetchone()[0]


_ledgers: dict[Path, Ledger] = {}


def open_ledger(path: Path) -> Ledger:
    path = path.resolve()
    if path not in _ledgers:
        _ledgers[path] = Ledger(path)
    return _ledgers[path]
